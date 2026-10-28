"""DOCX loader: headings, paragraphs, lists and tables in document order, plus core properties.

Unlike PDF, DOCX stores structure explicitly: heading and list styles, real tables. So this
loader reads structure instead of guessing it. Headers, footers and comments are not read.

Text formatted as hidden (Word's "Hidden" font effect) is dropped, like hidden HTML: Word
neither shows nor prints it, so it can only be meant for machines - a place to plant prompt
injections. Only direct formatting is checked, not hidden set through a character style.
"""

import io
import re
from datetime import datetime

from docx import Document
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph

from investigator.ingestion.models import Block, LoadedDocument, pick_title, table_blocks

_HEADING_STYLE = re.compile(r"Heading (\d)")


def load_docx(data: bytes, filename: str) -> LoadedDocument:
    doc = Document(io.BytesIO(data))
    blocks: list[Block] = []
    hidden_runs = 0

    def visible_text(paragraph: Paragraph) -> str:
        nonlocal hidden_runs
        parts = []
        for item in paragraph.iter_inner_content():  # runs and hyperlinks, in order
            for run in item.runs if isinstance(item, Hyperlink) else [item]:
                if run.font.hidden:
                    hidden_runs += 1
                else:
                    parts.append(run.text)
        return "".join(parts)

    # Paragraphs and tables interleaved, in the order they appear in the document.
    for item in doc.iter_inner_content():
        if isinstance(item, Table):
            rows = [
                ["\n".join(visible_text(p) for p in cell.paragraphs) for cell in row.cells]
                for row in item.rows
            ]
            blocks.extend(table_blocks(rows))
            continue
        text = visible_text(item).strip()
        if not text:
            continue
        style = item.style.name if item.style is not None else ""
        if heading := _HEADING_STYLE.match(style):
            blocks.append(Block("heading", text, level=int(heading.group(1))))
        elif style == "Title":
            blocks.append(Block("heading", text, level=1))
        elif style.startswith("List") or _has_numbering(item):
            blocks.append(Block("list_item", text))
        else:
            blocks.append(Block("paragraph", text))

    props = doc.core_properties
    metadata = {
        "author": props.author,
        "created_at": _iso(props.created),
        "modified_at": _iso(props.modified),
        "last_modified_by": props.last_modified_by,
        "doc_type": props.category,
        "keywords": props.keywords,
        "subject": props.subject,
        "revision": props.revision,
        "hidden_runs_removed": hidden_runs,
    }
    return LoadedDocument(pick_title(props.title, blocks, filename), "docx", blocks, metadata)


def _has_numbering(paragraph: Paragraph) -> bool:
    """Lists made with Word's bullet/number buttons carry numbering properties (numPr)."""
    properties = paragraph._p.pPr  # python-docx has no public API for this
    return properties is not None and properties.numPr is not None


def _iso(value: datetime | None) -> str | None:
    return value.replace(tzinfo=None).isoformat() if value else None
