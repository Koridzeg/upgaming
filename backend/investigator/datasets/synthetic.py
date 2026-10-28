"""Render the synthetic incident corpus into its final file formats.

Sources live in data/synthetic/source/. Markdown sources whose front matter says
`render: pdf` or `render: docx` are converted to that format, with title, author and date
written into the file's own metadata (the PDF info dictionary, DOCX core properties), so the
loaders have real metadata to extract. Every other source (.md, .html, .json, .csv) is already
in its final format and is copied unchanged. Output goes to data/synthetic/corpus/.

Usage: python -m investigator.datasets.synthetic
"""

import shutil
import zipfile
from datetime import UTC, datetime
from pathlib import Path

from investigator.config import DATA_DIR
from investigator.ingestion.markdown import (
    MarkdownBlock,
    parse_markdown,
    split_front_matter,
    strip_inline,
)

SYNTHETIC_DIR = DATA_DIR / "synthetic"
SOURCE_DIR = SYNTHETIC_DIR / "source"
CORPUS_DIR = SYNTHETIC_DIR / "corpus"

# A byte-identical copy under another name: the "same file uploaded twice" dedup case.
EXACT_COPIES = {"postmortem-INC-2026-0311 (copy).pdf": "postmortem-INC-2026-0311.pdf"}


plain = strip_inline  # PDF/DOCX would show inline Markdown markers literally


def _date(meta: dict[str, str]) -> datetime:
    return datetime.fromisoformat(meta["date"]).replace(tzinfo=UTC)


def render_pdf(meta: dict[str, str], blocks: list[MarkdownBlock], path: Path) -> None:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_title(meta["title"])
    pdf.set_author(meta["author"])
    pdf.set_subject(meta.get("doc_type", ""))
    pdf.set_keywords(f"status: {meta.get('status', '')}")
    pdf.set_creation_date(_date(meta))
    pdf.set_margins(20, 20, 20)
    pdf.add_page()

    for block in blocks:
        if block.kind == "heading":
            pdf.set_font("Helvetica", "B", {1: 16, 2: 13}.get(block.level, 11))
            pdf.ln(3)
            pdf.multi_cell(0, 8, plain(block.text), new_x="LMARGIN", new_y="NEXT")
        elif block.kind == "paragraph":
            pdf.set_font("Helvetica", size=10.5)
            pdf.multi_cell(0, 5.5, plain(block.text), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
        elif block.kind in ("bullets", "numbered"):
            pdf.set_font("Helvetica", size=10.5)
            for i, item in enumerate(block.items, start=1):
                marker = "-" if block.kind == "bullets" else f"{i}."
                pdf.multi_cell(0, 5.5, f"{marker} {plain(item)}", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
        elif block.kind == "table":
            pdf.set_font("Helvetica", size=9.5)
            with pdf.table(first_row_as_headings=True) as table:
                for row in block.items:
                    cells = table.row()
                    for cell in row:
                        cells.cell(plain(cell))
            pdf.ln(2)
    pdf.output(str(path))


def render_docx(meta: dict[str, str], blocks: list[MarkdownBlock], path: Path) -> None:
    from docx import Document

    doc = Document()
    props = doc.core_properties
    props.title = meta["title"]
    props.author = meta["author"]
    props.last_modified_by = meta["author"]
    props.created = props.modified = _date(meta)
    props.category = meta.get("doc_type", "")
    props.keywords = f"status: {meta.get('status', '')}"
    props.revision = 1

    for block in blocks:
        if block.kind == "heading":
            doc.add_heading(plain(block.text), level=min(block.level, 3))
        elif block.kind == "paragraph":
            for line in block.text.split("\n"):
                doc.add_paragraph(plain(line))
        elif block.kind in ("bullets", "numbered"):
            style = "List Bullet" if block.kind == "bullets" else "List Number"
            for item in block.items:
                doc.add_paragraph(plain(item), style=style)
        elif block.kind == "table":
            table = doc.add_table(rows=len(block.items), cols=len(block.items[0]))
            table.style = "Table Grid"
            for r, row in enumerate(block.items):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = plain(cell)
    doc.save(str(path))
    _fix_zip_timestamps(path)


def _fix_zip_timestamps(path: Path) -> None:
    """Make a DOCX byte-for-byte reproducible.

    A DOCX is a zip archive, and python-docx stamps every entry with the current time, so each
    build would produce different bytes (and a git diff) for identical content.
    """
    with zipfile.ZipFile(path) as source:
        entries = [(info.filename, source.read(info)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w") as target:
        for name, content in entries:
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            target.writestr(info, content, compress_type=zipfile.ZIP_DEFLATED)


RENDERERS = {"pdf": render_pdf, "docx": render_docx}


def build(source_dir: Path = SOURCE_DIR, corpus_dir: Path = CORPUS_DIR) -> list[Path]:
    """Regenerate the whole corpus directory from the sources."""
    if corpus_dir.exists():
        shutil.rmtree(corpus_dir)  # generated output only: never edit files in it by hand
    corpus_dir.mkdir(parents=True)

    written = []
    for source in sorted(source_dir.iterdir()):
        meta: dict[str, str] = {}
        if source.suffix == ".md":
            meta, body = split_front_matter(source.read_text(encoding="utf-8"))
        target = meta.get("render")
        if target:
            out = corpus_dir / f"{source.stem}.{target}"
            RENDERERS[target](meta, parse_markdown(body), out)
        else:
            out = corpus_dir / source.name
            shutil.copyfile(source, out)
        written.append(out)

    for copy_name, original in EXACT_COPIES.items():
        shutil.copyfile(corpus_dir / original, corpus_dir / copy_name)
        written.append(corpus_dir / copy_name)
    return written


def main() -> None:
    for path in build():
        print(f"{path.stat().st_size:>8,} B  {path.name}")


if __name__ == "__main__":
    main()
